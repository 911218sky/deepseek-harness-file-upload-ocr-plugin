#!/usr/bin/env python3
"""Extract common document formats and OCR PDF/image content locally."""

import argparse
import io
import json
import sys
from html.parser import HTMLParser
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from rapidocr_onnxruntime import RapidOCR


class TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


def truncate_output(output: str, max_chars: int) -> str:
    output = output.strip()
    if len(output) <= max_chars:
        return output
    notice = f"\n\n--- 已截断 / Truncated: output limited to {max_chars} characters ---"
    return output[: max(0, max_chars - len(notice))].rstrip() + notice


def joined_length(parts: list[str]) -> int:
    if not parts:
        return 0
    return sum(len(part) for part in parts) + 2 * (len(parts) - 1)


def ocr_render_scale(page_limit: int, dpi: int) -> float:
    """Lower render scale for long scanned PDFs so OCR stays within time/memory limits."""
    if page_limit <= 40:
        effective = dpi
    elif page_limit <= 100:
        effective = min(dpi, 120)
    else:
        effective = min(dpi, 96)
    return effective / 72


def pdf_text(data: bytes, args: argparse.Namespace) -> str:
    document = pdfium.PdfDocument(data)
    total_pages = len(document)
    page_limit = min(total_pages, args.max_pages)
    engine = None
    pages: list[str] = []
    render_scale = ocr_render_scale(page_limit, args.dpi)
    reserve = 120

    if total_pages > args.max_pages:
        pages.append(
            f"--- 已截断 / Truncated: processed first {page_limit} of {total_pages} pages "
            f"(configured limit {args.max_pages}) ---"
        )

    for index in range(page_limit):
        page = document[index]
        text_page = page.get_textpage()
        try:
            text = text_page.get_text_range().strip()
            if len(text) < args.native_text_min_chars:
                if engine is None:
                    engine = RapidOCR()
                image = page.render(scale=render_scale).to_pil()
                try:
                    encoded = io.BytesIO()
                    image.save(encoded, format="PNG")
                    result, _ = engine(encoded.getvalue())
                    text = "\n".join(item[1] for item in (result or []))
                finally:
                    image.close()
            page_block = f"--- 第 {index + 1} 页 / Page {index + 1} ---\n{text}"
            next_len = joined_length(pages) + len(page_block) + (2 if pages else 0)
            if next_len > args.max_output_chars - reserve:
                pages.append(
                    f"--- 已截断 / Truncated: stopped at page {index + 1} "
                    f"due to output limit ({args.max_output_chars} characters) ---"
                )
                break
            pages.append(page_block)
        finally:
            text_page.close()
            page.close()

    document.close()
    return "\n\n".join(pages)


def image_text(data: bytes) -> str:
    result, _ = RapidOCR()(data)
    return "\n".join(item[1] for item in (result or []))


def word_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(blocks)


def excel_text(data: bytes, max_sheets: int, max_output_chars: int) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        total_sheets = len(workbook.sheetnames)
        sheet_limit = min(total_sheets, max_sheets)
        sheets: list[str] = []
        reserve = 120

        if total_sheets > max_sheets:
            sheets.append(
                f"--- 已截断 / Truncated: processed first {sheet_limit} of {total_sheets} sheets "
                f"(configured limit {max_sheets}) ---"
            )

        for worksheet in workbook.worksheets[:sheet_limit]:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append("\t".join(values).rstrip())
            block = f"--- 工作表 / Sheet: {worksheet.title} ---\n" + "\n".join(rows)
            next_len = joined_length(sheets) + len(block) + (2 if sheets else 0)
            if next_len > max_output_chars - reserve:
                sheets.append(
                    f"--- 已截断 / Truncated: stopped at sheet {worksheet.title} "
                    f"due to output limit ({max_output_chars} characters) ---"
                )
                break
            sheets.append(block)
        return "\n\n".join(sheets)
    finally:
        workbook.close()


def powerpoint_text(data: bytes, max_slides: int, max_output_chars: int) -> str:
    presentation = Presentation(io.BytesIO(data))
    total_slides = len(presentation.slides)
    slide_limit = min(total_slides, max_slides)
    slides: list[str] = []
    reserve = 120

    if total_slides > max_slides:
        slides.append(
            f"--- 已截断 / Truncated: processed first {slide_limit} of {total_slides} slides "
            f"(configured limit {max_slides}) ---"
        )

    for index, slide in enumerate(list(presentation.slides)[:slide_limit]):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
        block = f"--- 第 {index + 1} 张幻灯片 / Slide {index + 1} ---\n" + "\n".join(parts)
        next_len = joined_length(slides) + len(block) + (2 if slides else 0)
        if next_len > max_output_chars - reserve:
            slides.append(
                f"--- 已截断 / Truncated: stopped at slide {index + 1} "
                f"due to output limit ({max_output_chars} characters) ---"
            )
            break
        slides.append(block)
    return "\n\n".join(slides)


def decoded_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("文本文件不是 UTF-8 或 GB18030 编码 / Text file is neither UTF-8 nor GB18030 encoded.")


def main() -> None:
    try:
        parser = argparse.ArgumentParser()
        parser.add_argument("--filename", required=True)
        parser.add_argument("--max-pages", type=int, required=True)
        parser.add_argument("--dpi", type=int, required=True)
        parser.add_argument("--native-text-min-chars", type=int, required=True)
        parser.add_argument("--max-output-chars", type=int, required=True)
        args = parser.parse_args()

        data = sys.stdin.buffer.read()
        if not data:
            raise ValueError("文件为空 / File upload is empty.")
        suffix = Path(args.filename).suffix.lower()
        if suffix == ".pdf":
            kind, output = "pdf", pdf_text(data, args)
        elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
            kind, output = "image", image_text(data)
        elif suffix == ".docx":
            kind, output = "word", word_text(data)
        elif suffix in {".xlsx", ".xlsm"}:
            kind, output = "excel", excel_text(data, args.max_pages, args.max_output_chars)
        elif suffix == ".pptx":
            kind, output = "powerpoint", powerpoint_text(data, args.max_pages, args.max_output_chars)
        elif suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".log", ".py", ".js", ".ts", ".tsx", ".css"}:
            kind, output = "text", decoded_text(data)
        elif suffix in {".html", ".htm"}:
            html = TextHTMLParser()
            html.feed(decoded_text(data))
            kind, output = "html", "\n".join(html.parts)
        elif suffix in {".doc", ".xls", ".ppt"}:
            raise ValueError(
                f"不支持旧版 Office 格式 {suffix}。请先另存为 "
                f"{suffix}x 后再上传 / Legacy Office format {suffix} is not supported. "
                f"Please save as {suffix}x and upload again."
            )
        else:
            raise ValueError(
                "不支持的文件类型。支持 PDF、图片、DOCX、XLSX/XLSM、PPTX 及常见文本/CSV/JSON/HTML 文件 / "
                "Unsupported file type. Supported: PDF, images, DOCX, XLSX/XLSM, PPTX, and common text/CSV/JSON/HTML files."
            )

        output = truncate_output(output, args.max_output_chars)
        print(json.dumps({"kind": kind, "text": output}, ensure_ascii=True))
    except Exception as error:
        print(str(error), file=sys.stderr)
        raise SystemExit(1) from error


if __name__ == "__main__":
    main()
